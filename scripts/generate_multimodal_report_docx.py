#!/usr/bin/env python3
"""Generate team meeting DOCX: multimodal interaction & product roadmap."""

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUTPUT = Path(__file__).resolve().parent.parent / "docs" / "组会汇报_多模态交互与产品形态规划.docx"


def set_run_font(run, name_cn="宋体", name_en="Times New Roman", size=11, bold=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = name_en
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name_cn)


def add_title(doc: Document, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run_font(run, size=18, bold=True)


def add_subtitle(doc: Document, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run_font(run, size=11)
    run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)


def add_heading(doc: Document, text: str, level: int = 1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, size=14 if level == 1 else 12, bold=True)


def add_para(doc: Document, text: str, bold: bool = False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, bold=bold)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.25
    return p


def add_bullet(doc: Document, text: str, level: int = 0):
    p = doc.add_paragraph(text, style="List Bullet")
    p.paragraph_format.left_indent = Cm(0.75 * (level + 1))
    for run in p.runs:
        set_run_font(run)


def add_table(doc: Document, headers: list[str], rows: list[list[str]]):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        for p in hdr_cells[i].paragraphs:
            for run in p.runs:
                set_run_font(run, bold=True, size=10)
    for r_idx, row in enumerate(rows):
        row_cells = table.rows[r_idx + 1].cells
        for c_idx, val in enumerate(row):
            row_cells[c_idx].text = val
            for p in row_cells[c_idx].paragraphs:
                for run in p.runs:
                    set_run_font(run, size=10)
    doc.add_paragraph()


def build_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.8)
    section.right_margin = Cm(2.8)

    add_title(doc, "AI 健身教练项目：多模态交互与产品形态规划")
    add_subtitle(doc, f"Adaptive Athletic Risk Assessment System · 组会汇报材料 · {date.today().isoformat()}")

    add_para(
        doc,
        "本文档基于项目当前进展（风险评估、安全计划生成、AI 教练对话、伤病档案同步重校验、"
        "前端可演示 MVP 及核心安全自动化测试），提出多模态交互演进路径与真实健身场景下的产品形态建议。",
    )

    # --- Section 1 ---
    add_heading(doc, "一、项目现状与多模态扩展思路", 1)
    add_para(
        doc,
        "当前系统采用纯文本交互链路：用户发送文本消息 → build_context 注入用户档案、伤病史、"
        "训练历史与活跃计划 → CoachAgent 经 LLM 生成回复（支持 SSE 流式输出）。"
        "多模态扩展无需推翻现有架构，建议在既有链路上增加「模态适配层」，将各类输入统一转换为"
        "语义等价的文本上下文后，继续走 CoachAgent 与 safety_sync 安全闭环。",
    )
    add_para(doc, "建议架构：", bold=True)
    add_bullet(doc, "文字输入：沿用现有 /chat/stream，强化快捷话术与组间一句话记录")
    add_bullet(doc, "语音输入：STT（语音转文字）→ 现有 ChatRequest.message")
    add_bullet(doc, "图片/视频：Vision 模型解析 → 摘要为文本 → 注入 CoachAgent 上下文")
    add_bullet(doc, "语音输出：LLM 文本回复 → TTS（文字转语音）→ 客户端流式播放")

    # --- Section 2 ---
    add_heading(doc, "二、多模态能力分期规划", 1)

    add_heading(doc, "2.1 文字交互（已有基础，持续强化）", 2)
    add_bullet(doc, "适用场景：计划讨论、训练复盘、改计划确认")
    add_bullet(doc, "示例：「深蹲第三组自觉强度 8，左膝有点紧」")
    add_bullet(doc, "优势：与现有 action_request 机制天然兼容，开发成本最低")

    add_heading(doc, "2.2 语音交互（健身场景最高优先级 · P0）", 2)
    add_para(doc, "用户侧：浏览器 MediaRecorder / 小程序录音 / App 原生录音 → 上传音频")
    add_para(doc, "服务端：Whisper 或国内 STT 服务 → 转写为文本 → 调用现有教练 API")
    add_para(doc, "教练侧：LLM 文本回复 → TTS 合成 → 流式播报")
    add_para(doc, "关键产品设计：", bold=True)
    add_bullet(doc, "训练模式：默认语音入、语音出，界面仅展示组数、休息倒计时、禁忌提醒等要点")
    add_bullet(doc, "接口扩展：在现有 SSE delta/final 事件外，增加 audio_chunk 或 final 后异步 /tts")
    add_bullet(doc, "体验优化：长回复分段合成，首句优先播报以降低感知延迟")

    add_heading(doc, "2.3 图片与视频（差异化能力 · P1/P2）", 2)
    add_table(
        doc,
        ["场景", "输入形式", "处理方式"],
        [
            ["动作纠错", "深蹲/硬拉短视频", "Vision 模型 + 动作库对照，结合伤病史给出替代建议"],
            ["疼痛反馈", "肿胀、淤青照片", "不作医学诊断，记录症状并触发 safety_sync 重校验提示"],
            ["器械环境", "健身房环境照片", "识别可用器械，动态调整当日训练内容"],
            ["数据截图", "手表/体脂秤截图", "OCR 提取数据并写入训练档案"],
        ],
    )
    add_para(
        doc,
        "技术路径：扩展 ChatRequest 为 content_parts（text / image_url / audio_url），"
        "Vision 结果摘要为文本后再进入 CoachAgent，与现有 build_context 保持兼容。",
    )

    add_heading(doc, "2.4 分期优先级汇总", 2)
    add_table(
        doc,
        ["阶段", "能力", "理由"],
        [
            ["P0", "语音输入 + 语音输出", "健身房双手占用、出汗，打字体验差，语音最贴合组间场景"],
            ["P1", "拍照记录疼痛/训练日志", "强化伤病安全闭环，便于实验评估与论文撰写"],
            ["P2", "短视频动作分析", "算力与延迟成本高，更适合原生 App 相机能力"],
        ],
    )

    # --- Section 3 ---
    add_heading(doc, "三、网站形态在真实健身场景的局限", 1)
    add_table(
        doc,
        ["问题", "说明"],
        [
            ["双手占用", "组间休息仅 30–90 秒，打字操作成本高"],
            ["后台与锁屏", "浏览器易被系统回收，计时器与语音播报易中断"],
            ["硬件能力受限", "麦克风权限繁琐；难以稳定对接蓝牙心率、Apple Health 等"],
            ["网络依赖", "地下室健身房信号弱，纯在线 Web 体验不稳定"],
            ["使用心智不匹配", "用户较少在训练时主动「打开网站」，更像管理工具而非训练伴侣"],
        ],
    )
    add_para(
        doc,
        "结论：Web 适合承担「教练办公室」职能——计划制定、档案管理、风险分析、数据复盘；"
        "不适合承担「健身房实时陪练」职能——组间指导、语音互动、可穿戴联动。",
    )

    # --- Section 4 ---
    add_heading(doc, "四、产品形态对比：Web / 小程序 / App", 1)
    add_table(
        doc,
        ["维度", "网站", "小程序", "原生 App"],
        [
            ["获客成本", "低", "很低（分享、搜一搜）", "高"],
            ["开发成本", "已完成 MVP", "中等（复用 API 与部分逻辑）", "高"],
            ["语音/相机", "一般", "较好", "最好"],
            ["后台/锁屏", "差", "一般", "好"],
            ["支付与会员", "一般", "成熟", "成熟"],
            ["演示与论文", "适合", "适合", "适合"],
            ["可穿戴/HealthKit", "几乎无", "有限", "完整"],
        ],
    )

    add_heading(doc, "4.1 落地性判断", 2)
    add_para(doc, "短期最落地：微信小程序（或支付宝小程序）", bold=True)
    add_para(doc, "中期差异化：React Native / Flutter App（或 PWA + 原生壳）", bold=True)
    add_para(doc, "长期品牌与数据资产：独立原生 App", bold=True)
    add_para(
        doc,
        "推荐策略并非二选一，而是「小程序获客验证 → App 做深度体验」的分阶段演进。",
    )

    add_heading(doc, "4.2 小程序更适合当前阶段的原因", 2)
    add_bullet(doc, "后端改动小：FastAPI + SSE 可直接服务小程序，多模态仅需增补 STT/TTS/上传接口")
    add_bullet(doc, "用户习惯匹配：国内健身用户微信使用率高，计划分享与打卡传播成本低")
    add_bullet(doc, "场景覆盖足够：组间语音问答、疼痛拍照记录、今日计划查看均可实现")
    add_bullet(doc, "交付周期短：Taro / uni-app 可复用部分 React 思路，快于双端原生开发")

    add_heading(doc, "4.3 何时需要独立 App", 2)
    add_bullet(doc, "Apple Watch / 心率带 / 蓝牙器械实时联动")
    add_bullet(doc, "弱网环境下的离线计划与本地 TTS")
    add_bullet(doc, "长时间后台语音教练（类似跟练类产品）")
    add_bullet(doc, "订阅制品牌运营，独立 icon、推送与数据沉淀")

    # --- Section 5 ---
    add_heading(doc, "五、推荐演进路线", 1)
    add_para(doc, "Phase A（约 1–2 周）：Web 增加「训练模式」", bold=True)
    add_bullet(doc, "大按钮语音输入 + TTS 播报 + 简化 UI")
    add_bullet(doc, "复用现有 /chat/stream 流式接口")

    add_para(doc, "Phase B（约 3–4 周）：微信小程序 MVP", bold=True)
    add_bullet(doc, "核心功能：今日计划、语音教练、练后打卡、伤病同步")
    add_bullet(doc, "后端增补：/stt、/tts、multipart 图片消息接口")

    add_para(doc, "Phase C（验证后）：App 或 RN 一套代码", bold=True)
    add_bullet(doc, "离线缓存、推送提醒、HealthKit、动作视频分析")

    add_para(
        doc,
        "产品定位一句话：Web 做「教练办公室」（计划、档案、分析）；"
        "小程序/App 做「健身房里的教练」（语音、快捷记录、组间指导）。",
        bold=True,
    )

    # --- Section 6 ---
    add_heading(doc, "六、研究与论文结合点（可选）", 1)
    add_bullet(doc, "输入模态：文本 vs 语音在组间场景下的可用性与任务完成时间对比")
    add_bullet(doc, "视觉模态：疼痛自述 + 图片是否提升 safety_sync 触发率与计划调整率")
    add_bullet(doc, "输出模态：TTS 教练 vs 纯文字对训练依从性（完成率、RPE 记录率）的影响")
    add_para(
        doc,
        "上述实验设计可将工程落地与可量化研究贡献结合，适合作为后续 Sprint 或论文扩展章节。",
    )

    # --- Milestone ---
    add_heading(doc, "附录：当前项目 Milestone（英文）", 1)
    add_para(
        doc,
        "Delivered an injury-aware training MVP with end-to-end coverage from risk assessment "
        "and safe plan generation to AI coaching and profile sync/revalidation, a demo-ready "
        "frontend, and automated tests for core safety logic.",
    )

    return doc


def main():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = build_document()
    doc.save(OUTPUT)
    print(f"Saved: {OUTPUT}")


if __name__ == "__main__":
    main()
